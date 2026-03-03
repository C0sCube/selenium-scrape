import base64, tempfile, pdfkit,json, pandas as pd, os, pdfplumber, re
from io import BytesIO
from PyPDF2 import PdfMerger
from datetime import datetime
from app.constants import get_paths

import pandas as pd
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
class PDFReportBuilderPro:
    def __init__(self, cache_data = None, output_path = None):
        self.cache_data = cache_data
        self.output_path = output_path
        self.merger = PdfMerger()
        self.bank_titles = []
        self.config = pdfkit.configuration(
            wkhtmltopdf=get_paths()["htmltopdf_path"]
        )

        # color palette for per-bank tinting (soft pastels)
        self.color_palette = [
            "#f9f9ff", "#fff9f9", "#f9fff9", "#fffaf2", "#f2fbff", "#fff7fc"
        ]

        # shared pdfkit footer setup
        self.pdf_options = {
            "footer-right": "Page [page] of [toPage]",
            "footer-font-size": "9",
            "footer-spacing": "5",
            "margin-top": "15mm",
            "margin-bottom": "15mm",
            "margin-left": "10mm",
            "margin-right": "10mm"
        }

    # ---------- helpers ----------
    def _make_bank_header(self, bank_name, summary_text, color):
        """Top banner for each bank section."""
        return f"""
        <div style='background:{color};padding:12px 16px;margin-bottom:10px;
                    border-radius:6px;border-left:6px solid #1f4e79;'>
            <h1 style='font-size:18pt;color:#1f4e79;margin-bottom:4px;'>{bank_name}</h1>
            <p style='font-size:10pt;color:#333;margin:0;'>{summary_text}</p>
        </div>
        """

    def _make_action_header(self, scrape):
        
        unattatch_actions = ["website"]
        """Header block before each scrape section."""
        action = scrape.get("action", "")
        if action in unattatch_actions: return None
        timestamp = scrape.get("timestamp", "")
        data_present = scrape.get("data_present", "")
        count = scrape.get("response_count", 0) if data_present else 0
        webpage = scrape.get("webpage", "")
        return f"""
        <div style='margin:6px 0;padding:8px 10px;border-left:4px solid #1f4e79;
                    background:#f8fafc;border-radius:4px;'>
            <h3 style='color:#1f4e79;margin-bottom:2px;'>Action: {action}</h3>
            <p style='font-size:9pt;color:#444;margin:0;'>
                Timestamp: {timestamp} | Present: {data_present} | Count: {count}<br>
                Website: <a href='{webpage}' target='_blank'>{webpage}</a>
            </p>
        </div>
        """

    def _make_title_block(self, titles):
        """Format scraped titles."""
        if not titles:
            return ""
        html_titles = ""
        for idx, t in enumerate(titles):
            html_titles += f"<p style='font-size:9pt;color:#222;margin:0;'>• {t}</p>"
        return f"<div style='margin:4px 0 6px 10px;'>{html_titles}</div>"

    def _make_table_separator(self):
        """Separator between tables or actions."""
        return """
        <div style='text-align:center;color:#999;font-size:8.5pt;
                    margin:12px 0;'>──── xxx ────</div>
        """
        
    def _html_to_pdf(self, html_content, output_file):
        """Convert HTML string to PDF using pdfkit."""
        html_template = f"""
        <html>
        <head>
            <meta charset="utf-8">
            <style>
                @page {{
                    size: A4;
                    margin: 1.5cm 1cm 1.5cm 1cm;
                }}
                body {{
                    font-family: Helvetica, Arial, sans-serif;
                    font-size: 9pt;
                    line-height: 1.25;
                    color: #111;
                }}
                h1,h2,h3 {{ color: #1f4e79; }}
                table {{
                    border-collapse: collapse;
                    width: 95%;
                    margin: 6px auto;
                }}
                th,td {{
                    border: 1px solid #999;
                    padding: 3px 5px;
                    font-size: 8.5pt;
                }}
                th {{
                    background-color: #f0f0f0;
                }}
                tr:nth-child(even) {{
                    background-color: #fafafa;
                }}
                tr:hover {{
                    background-color: #eaf2fb;
                }}
                p {{ margin: 2px 0; }}
                a {{ color: #1f4e79; text-decoration: none; }}

                /* ===== API Response Styling ===== */
                .api-block {{
                    background: #f5f9ff;
                    border-left: 4px solid #0078d7;
                    border-radius: 6px;
                    padding: 8px 10px;
                    margin: 10px 0;
                }}
                .api-header {{
                    font-weight: bold;
                    color: #0f4c81;
                    font-size: 9pt;
                    margin-bottom: 6px;
                }}
                .api-block h4 {{
                    color: #0f4c81;
                    margin-top: 6px;
                    font-size: 9pt;
                    text-transform: uppercase;
                }}
                .api-block table {{
                    width: 94%;
                    margin: 4px auto;
                }}
                .api-block p {{
                    margin: 2px 0;
                    font-size: 9pt;
                }}
                .api-block .api-url {{
                    color: #555;
                    font-size: 8pt;
                    margin-top: 4px;
                }}
            </style>
        </head>
        <body>
            {html_content}
        </body>
        </html>
        """
        pdfkit.from_string(
            html_template,
            output_file,
            configuration=self.config,
            options=self.pdf_options
        )

    # ---------- main build ----------
    def build(self,cache_json,path):
        toc_rows = []
        temp_files = []
        self.cache_data= cache_json
        self.output_path = path

        # ===== FRONT PAGE =====
        cover_html = f"""
        <div style='text-align:center;margin-top:180px;'>
            <h1 style='font-size:28pt;color:#1f4e79;'>Deposit Rate Report</h1>
            <p style='font-size:12pt;color:#444;'>Generated on {datetime.now().strftime("%d %b %Y, %H:%M")}</p>
            <hr style='width:60%;border:1px solid #1f4e79;margin:20px auto;'>
        </div>
        <div style='margin:40px auto;width:70%;font-size:11pt;color:#222;'>
            <h2 style='text-align:center;color:#1f4e79;'>Table of Contents</h2>
            <table style='width:100%;border-collapse:collapse;margin-top:10px;'>
                <thead style='background:#f0f0f0;'>
                    <tr>
                        <th style='border:1px solid #ccc;padding:5px;'>#</th>
                        <th style='border:1px solid #ccc;padding:5px;'>Bank Code</th>
                        <th style='border:1px solid #ccc;padding:5px;'>Bank Name</th>
                    </tr>
                </thead>
                <tbody>
        """

        # ===== Build each bank section =====
        for i, record in enumerate(self.cache_data.get("records", []), start=1):
            bank_name = record.get("bank_name", "Unknown Bank")
            bank_code = record.get("bank_code", "N/A")
            color = self.color_palette[(i - 1) % len(self.color_palette)]

            self.bank_titles.append(bank_name)

            summary_text = (f"Compiled on {datetime.now().strftime('%d %b %Y')}"  f"from {len(record.get('scraped_data', []))} sources.")

            html_parts = [self._make_bank_header(bank_name, summary_text, color)]
            pdf_files = []
            added_anything = False

            # collect scrape data
            for scrape in record.get("scraped_data", []):
                
                action_header = self._make_action_header(scrape)
                if action_header:
                    html_parts.append(action_header)
                    
                responses = scrape.get("response", [])
                for response in responses:
                    #value
                    typ = response.get("type")
                    value = response.get("value", "")
                    if not value:
                        continue
                    
                    #title
                    titles = response.get("title", [])
                    if len(titles): #if only
                        html_parts.append(self._make_title_block(titles))

                    #format to save report
                    if typ in ("html", "table_html"):
                        html_str = value.strip()
                        if len(html_str) < 150:
                            continue
                        html_parts.append(html_str)
                        html_parts.append(self._make_table_separator())
                        added_anything = True

                    elif typ == "xlsx":
                        try:
                            xlsx_bytes = base64.b64decode(value)
                            df = pd.read_excel(BytesIO(xlsx_bytes))
                            if df.empty:
                                continue
                            html_parts.append(df.to_html(index=False))
                            html_parts.append(self._make_table_separator())
                            added_anything = True
                        except Exception:
                            continue

                    elif typ == "pdf":
                        try:
                            pdf_bytes = base64.b64decode(value)
                            tmp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
                            tmp_pdf.write(pdf_bytes)
                            tmp_pdf.close()
                            pdf_files.append(tmp_pdf.name)
                            added_anything = True
                        except Exception:
                            continue
                        
                    elif typ == "json":
                        try:
                            value = response.get("value")
                            structure = response.get("structure", {})
                            sections = self.parse_structured_json_for_report(value, structure)

                            html_parts.append("<div class='api-block'>")
                            html_parts.append("<div class='api-header'>📡 API Response</div>")

                            # Optional: show source URL if present
                            if "api_url" in response:
                                html_parts.append(f"<p class='api-url'>Source: {response['api_url']}</p>")

                            for key, section in sections:
                                if isinstance(section, str):
                                    html_parts.append(section)
                                elif hasattr(section, "to_html"):
                                    # add small visual polish
                                    html_parts.append(f"<h4 style='margin-top:8px;color:#1f4e79;'>{key}</h4>")
                                    if section.empty:
                                        html_parts.append(f"<p style='font-size:8.5pt;color:#888;'>No data for {key}</p>")
                                    else:
                                        html_parts.append(section.to_html(index=False))
                                html_parts.append("<br>")

                            html_parts.append("</div>")  # end .api-block
                            html_parts.append(self._make_table_separator())
                            added_anything = True

                        except Exception as e:
                            html_parts.append(f"""
                            <div class="api-block">
                                <p style="color:red;font-size:9pt;">
                                    ⚠️ Error rendering JSON: {str(e)}
                                </p>
                            </div>
                            """)
                            html_parts.append(self._make_table_separator())
                            added_anything = True

            if added_anything:
                combined_html = "".join(html_parts)
                tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
                self._html_to_pdf(combined_html, tmp.name)
                self.merger.append(tmp.name)
                temp_files.append(tmp.name)
                for pdf in pdf_files:
                    self.merger.append(pdf)

                toc_rows.append(f"""
                    <tr>
                        <td style='border:1px solid #ccc;padding:4px;text-align:center;'>{i}</td>
                        <td style='border:1px solid #ccc;padding:4px;text-align:center;'>{bank_code}</td>
                        <td style='border:1px solid #ccc;padding:4px;'>{bank_name}</td>
                    </tr>
                """)

        # finalize TOC
        cover_html += "".join(toc_rows) + "</tbody></table></div>"
        cover_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        self._html_to_pdf(cover_html, cover_pdf.name)
        self.merger.merge(position=0, fileobj=cover_pdf.name)
        cover_pdf.close()

        # ===== finalize merged pdf =====
        self.merger.write(self.output_path)
        self.merger.close()

        print(f"✅ Final PDF report generated at {self.output_path}")


    def parse_structured_json_for_report(self, data: dict, structure: dict):
        if not data or not structure:
            return [("⚠ No Data", pd.DataFrame([["No data found in API response."]]))]
        if structure: data = {k: v for k, v in data.items() if k in structure}

        sections = []
        normalized_structure = []

        # Normalize structure definitions
        for key, rule in structure.items():
            parts = rule.split("|")
            val_type, order, interpret_as = parts
            order = int(order)
            normalized_structure.append((key, val_type, order, interpret_as))

        # Process each key
        for key, val_type, order, interpret_as in sorted(normalized_structure, key=lambda x: x[2]):
            if key not in data:
                continue
            val = data[key]

            # --- Case: Null Response placeholder ---
            if (
                key == "data"
                and isinstance(val, list)
                and len(val) == 1
                and isinstance(val[0], dict)
                and val[0].get("symbol") == "Null Response from API"
            ):
                html_block = """
                <p style='font-size:9pt;margin:2px 0;color:#777;'>
                    ⚠️ <b>No usable data returned from API.</b>
                </p>
                """
                sections.append((key, html_block))
                continue

            # --- Handle string values ---
            if val_type == "str":
                html_block = f"""
                <p style='font-size:9pt;margin:2px 0;'>
                    <b>{key}:</b> {val if val not in [None, '', 'null'] else '—'}
                </p>
                """
                sections.append((key, html_block))
                continue

            # --- Handle dicts (two-column tables) ---
            if isinstance(val, dict):
                df = pd.DataFrame(list(val.items()), columns=["Field", "Value"])
                sections.append((key, df))
                continue

            # --- Handle lists (multi-row tables) ---
            if isinstance(val, list):
                if val and isinstance(val[0], dict):
                    df = pd.json_normalize(val, sep=".")
                elif val:
                    df = pd.DataFrame({key: val})
                else:
                    df = pd.DataFrame([[f"No data found for '{key}'"]], columns=[key])
                sections.append((key, df))
                continue

            # --- Fallback (single value) ---
            df = pd.DataFrame([[str(val)]], columns=[key])
            sections.append((key, df))

        return sections

    def parse_structured_json_for_excel(self, data: dict, structure: dict):
        if not data or not structure:
            return [("⚠ No Data", "No data found in API response.")]
        
        if structure:
            data = {k: v for k, v in data.items() if k in structure}


        sections = []
        normalized = []

        # Normalize rule set
        for key, rule in structure.items():
            parts = rule.split("|")
            val_type, order, interpret_as = parts
            normalized.append((key, val_type.strip(), int(order), interpret_as.strip()))

        for key, val_type, order, interpret_as in sorted(normalized, key=lambda x: x[2]):
            if key not in data:
                continue
            val = data[key]

            # 🟡 Case 1: Null API Response
            if (
                key == "data"
                and isinstance(val, list)
                and len(val) == 1
                and isinstance(val[0], dict)
                and str(val[0].get("symbol", "")).lower() == "null response from api"
            ):
                sections.append((key, "⚠ No usable data returned from API"))
                continue
            
            if val_type == "str":
                text_value = val if val not in [None, "", "null"] else "—"
                sections.append((key, f"{key}: {text_value}"))
                continue

            if isinstance(val, dict):
                df = pd.DataFrame(list(val.items()), columns=["Field", "Value"])
                sections.append((key, df))
                continue

            # 🟣 Case 4: List (API dataset → multi-row table)
            if isinstance(val, list):
                if val and isinstance(val[0], dict):
                    df = pd.json_normalize(val, sep=".")  # flatten nested
                    sections.append((key, df))
                elif val:
                    df = pd.DataFrame({key: val})  # list of scalars
                    sections.append((key, df))
                else:
                    sections.append((key, f"No data found for '{key}'"))
                continue
            
            sections.append((key, str(val)))

        return sections

    def write_excel_report(self, cache_json, path):
        """Generate a detailed Excel report mirroring the PDF structure."""
        self.cache_data = cache_json
        self.output_path = path

        wb = Workbook()
        ws_summary = wb.active
        ws_summary.title = "Summary"
        
        def sanitize_Win_filename(name):
            # Replace illegal Windows characters with underscores
            return re.sub(r'[<>:"/\\|?*]', '_', name)

        # Style helpers
        header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
        border = Border(left=Side(style="thin"), right=Side(style="thin"),
                        top=Side(style="thin"), bottom=Side(style="thin"))
        align_center = Alignment(horizontal="center", vertical="center")

        # ===== SUMMARY HEADERS =====
        ws_summary.append(["#", "Bank Code", "Bank Name", "Type", "Sources Found"])
        for cell in ws_summary[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = border

        summary_row = 2

        # ===== PER BANK SHEETS =====
        for i, record in enumerate(self.cache_data.get("records", []), start=1):
            name = record.get("bank_name", "N/A")
            code = record.get("bank_code", "N/A")
            type_ = record.get("bank_type", "N/A")

            ws = wb.create_sheet(title=f"{name[:28]}_{i}"[:31]) # Excel sheet name limit
            # ws.append(["Bank Name", name])
            # ws.append(["Bank Code", code])
            # ws.append(["Bank Type", type_])
            # ws.append([])
            ws.append(["Domain_Details",f"{code}",f"{name}", f"{type_}"])
            ws.append([])
            
            scraped_data = record.get("scraped_data", [])
            total_sources = len(scraped_data)
            row_cursor = 3

            for scrape in scraped_data:
                
                ws.append(["Action_Details",f"{scrape.get('action','')}",f"{scrape.get('timestamp','')}", f"{scrape.get('webpage','')}"])
                ws.append([])
                # ws.append([f"Action: {scrape.get('action','')}"])
                # ws.append([f"Timestamp: {scrape.get('timestamp','')}"])
                # ws.append([f"Webpage: {scrape.get('webpage','')}"])
                # ws.append([])
                row_cursor += 2

                responses = scrape.get("response", [])
                for response in responses:
                    typ = response.get("type")
                    value = response.get("value", "")
                    if not value:
                        continue

                    titles = response.get("title", [])
                    titles = [titles] if isinstance(titles, str) else titles or []
                    if titles:
                        ws.append(["Title: " + ", ".join(titles)])

                    # ========== TYPE HANDLING ==========
                    try:
                        if typ in ("html", "table_html"):
                            df = pd.read_html(value)[0]
                            for r in dataframe_to_rows(df, index=False, header=True):
                                ws.append(r)

                        elif typ == "xlsx":
                            xlsx_bytes = base64.b64decode(value)
                            df = pd.read_excel(BytesIO(xlsx_bytes))
                            for r in dataframe_to_rows(df, index=False, header=True):
                                ws.append(r)

                        # elif typ == "json":
                        #     json_bytes = base64.b64decode(value)
                        #     json_text = json_bytes.decode("utf-8", errors="ignore")
                        #     data = json.loads(json_text)
                        #     df = pd.json_normalize(data)
                        #     for r in dataframe_to_rows(df, index=False, header=True):
                        #         ws.append(r)
                        
                        elif typ == "json":
                            value = response.get("value")
                            structure = response.get("structure", {})
                            sections = self.parse_structured_json_for_excel(value, structure)
                            if "api_url" in response:  ws.append([f"Source: {response['api_url']}"])
                            
                            for key, section in sections:
                                if isinstance(section, str):
                                    ws.append([section])
                                elif hasattr(section, "to_numpy"):
                                    df = section
                                    for r in dataframe_to_rows(df, index=False, header=True):
                                        ws.append(r)
                                else:
                                    ws.append(["⚠ Unknown data type"])

                        elif typ == "pdf":
                            try:
                                # Decode and save the PDF locally
                                pdf_bytes = base64.b64decode(value)

                                # Build attachment folder & filename
                                safe_name = sanitize_Win_filename("_".join(titles)) or f"attachment_{i}"
                                pdf_name = f"{safe_name}.pdf"
                                pdf_folder = os.path.join(os.path.dirname(path), "attachments")
                                os.makedirs(pdf_folder, exist_ok=True)
                                pdf_path = os.path.join(pdf_folder, pdf_name)

                                with open(pdf_path, "wb") as f:
                                    f.write(pdf_bytes)

                                # Extract tables from the saved PDF using pdfplumber
                                with pdfplumber.open(pdf_path) as pdf:
                                    for page in pdf.pages:
                                        tables = page.extract_tables()
                                        if not tables:
                                            ws.append([f"No tables found in {pdf_name} (Page {page.page_number})"])
                                            continue
                                        for table in tables:
                                            df = pd.DataFrame(table)
                                            ws.append([f"📎 PDF Table — {pdf_name} (Page {page.page_number})"])
                                            for row in dataframe_to_rows(df, index=False, header=False):
                                                ws.append(row)
                                            ws.append([])

                                # Add hyperlink to file
                                link_cell = ws.cell(row=ws.max_row + 1, column=1, value=f"Open PDF → {pdf_name}")
                                link_cell.hyperlink = pdf_path
                                link_cell.style = "Hyperlink"

                            except Exception as e:
                                ws.append([f"[Error handling PDF] {e}"])
                        else:
                            ws.append([f"Unsupported type: {typ}"])

                    except Exception as e:
                        ws.append([f"[Error parsing {typ}] {e}"])

                    ws.append([])  # add spacing between responses

            # ===== Append to Summary =====
            ws_summary.append([i, code, name,type_, total_sources])
            summary_row += 1

        # ===== Apply styling to Summary =====
        for row in ws_summary.iter_rows(min_row=2, max_row=summary_row, min_col=1, max_col=5):
            for cell in row:
                cell.alignment = align_center
                cell.border = border

        # Auto width
        for column_cells in ws_summary.columns:
            length = max(len(str(cell.value)) if cell.value else 0 for cell in column_cells)
            ws_summary.column_dimensions[column_cells[0].column_letter].width = length + 4

        # ===== Save file =====
        wb.save(path)
        print(f"✅ Excel report generated at {path}")




# import os, base64, tempfile, pdfkit, pandas as pd
# from io import BytesIO
# from PyPDF2 import PdfMerger, PdfWriter, PdfReader
# from datetime import datetime
# from pypdf import PdfReader, PdfWriter


# class PDFReportBuilderPro:
#     def __init__(self, cache_data, output_path):
#         self.cache_data = cache_data
#         self.output_path = output_path
#         self.merger = PdfMerger()
#         self.bank_titles = []
#         self.config = pdfkit.configuration(
#             wkhtmltopdf=r"C:\Program Files\wkhtmltopdf\bin\wkhtmltopdf.exe"
#         )

#     # ---------- helpers ----------
#     def _make_bank_header(self, bank_name, summary_text):
#         anchor = bank_name.replace(" ", "_")
#         return f"""
#         <div style='page-break-before: always;' id='{anchor}'>
#             <h1 style='font-size:18pt;color:#1f4e79;text-align:center;margin-bottom:6px;'>{bank_name}</h1>
#             <p style='font-size:10pt;text-align:center;color:#555;'>{summary_text}</p>
#             <hr style='border:none;border-top:2px solid #1f4e79;width:90%;margin:6px auto;'>
#         </div>
#         """

#     def _make_action_header(self, scrape):
#         action = scrape.get("action", "")
#         timestamp = scrape.get("timestamp", "")
#         data_present = scrape.get("data_present", "")
#         count = scrape.get("response_count", 0) if data_present else 0
#         webpage = scrape.get("webpage", "")
#         return f"""
#         <div style='margin-top:6px;'>
#             <h3 style='color:#1f4e79;'>Action: {action}</h3>
#             <p style='font-size:9pt;color:#444;'>
#                 Timestamp: {timestamp} | Present: {data_present} | Count: {count}<br>
#                 Website: <a href='{webpage}' target='_blank'>{webpage}</a>
#             </p>
#             <hr style='border:none;border-top:1px solid #bbb;width:95%;margin:4px auto;'>
#         </div>
#         """

#     def _make_title_block(self, titles):
#         if not titles:
#             return ""
#         html_titles = ""
#         for idx, t in enumerate(titles):
#             html_titles += f"<p style='font-size:9pt;color:#222;'>TITLE {idx+1}: {t}</p>"
#         return html_titles

#     def _make_table_separator(self):
#         return """
#         <br><br>
#         <div style='text-align:center;color:#888;font-size:9pt;margin:6px 0;'>
#             ───── End of Section ─────
#         </div>
#         <br><br>
#         """

#     def _html_to_pdf(self, html_content, output_file):
#         html_template = f"""
#         <html>
#         <head>
#             <meta charset="utf-8">
#             <style>
#                 @page {{ size: A4; margin: 0.7cm; }}
#                 body  {{ font-family: Helvetica, Arial, sans-serif; font-size: 9pt; line-height: 1.15; color: #111; }}
#                 h1,h2,h3 {{ color: #1f4e79; margin-bottom: 4px; }}
#                 table {{ border-collapse: collapse; width: 95%; margin: 6px auto; }}
#                 th,td {{ border: 1px solid #999; padding: 3px 5px; font-size: 8.5pt; }}
#                 th    {{ background-color: #f0f0f0; }}
#                 p     {{ margin: 2px 0; }}
#                 a     {{ color: #1f4e79; text-decoration: none; }}
#             </style>
#         </head>
#         <body>
#             {html_content}
#         </body>
#         </html>
#         """
#         pdfkit.from_string(html_template, output_file, configuration=self.config)

#     # ---------- main build ----------
#     def build(self):
#         toc_anchor = '<a name="Table_of_Contents"></a>'
#         toc_rows = []
#         temp_files = []

#         # ===== Build each bank section =====
#         for i, record in enumerate(self.cache_data.get("records", []), start=1):
#             bank_name = record.get("bank_name", "Unknown Bank")
#             bank_code = record.get("bank_code", "N/A")
#             bank_type = record.get("bank_type", "Commercial Bank")
#             bank_anchor = bank_name.replace(" ", "_")

#             self.bank_titles.append(bank_name)

#             summary_text = (
#                 f"Compiled on {datetime.now().strftime('%d %b %Y')} "
#                 f"from {len(record.get('scraped_data', []))} sources."
#             )

#             html_parts = [self._make_bank_header(bank_name, summary_text)]
#             pdf_files = []
#             added_anything = False

#             # collect data
#             for scrape in record.get("scraped_data", []):
#                 html_parts.append(self._make_action_header(scrape))
#                 responses = scrape.get("response", [])
#                 for response in responses:
#                     typ = response.get("type")
#                     value = response.get("value", "")
#                     if not value:
#                         continue

#                     titles = response.get("title", [])
#                     titles = [titles] if isinstance(titles, str) else titles
#                     html_parts.append(self._make_title_block(titles))

#                     if typ in ("html", "table_html"):
#                         html_str = value.strip()
#                         if len(html_str) < 150:
#                             continue
#                         html_parts.append(html_str)
#                         html_parts.append(self._make_table_separator())
#                         added_anything = True

#                     elif typ == "xlsx":
#                         try:
#                             xlsx_bytes = base64.b64decode(value)
#                             df = pd.read_excel(BytesIO(xlsx_bytes))
#                             if df.empty:
#                                 continue
#                             html_parts.append(df.to_html(index=False))
#                             html_parts.append(self._make_table_separator())
#                             added_anything = True
#                         except Exception:
#                             continue

#                     elif typ == "pdf":
#                         try:
#                             pdf_bytes = base64.b64decode(value)
#                             tmp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
#                             tmp_pdf.write(pdf_bytes)
#                             tmp_pdf.close()
#                             pdf_files.append(tmp_pdf.name)
#                             added_anything = True
#                         except Exception:
#                             continue

#             # ---- Add footer links ----
#             footer_links = f"""
#             <hr style='border:none;border-top:1px solid #bbb;width:90%;margin:8px auto;'>
#             <p style='text-align:center;font-size:8pt;color:#555;'>
#                 <a href='#{bank_anchor}'>⬆ Go to Bank Header</a>  |  
#                 <a href='#Table_of_Contents'>⬅ Back to Table of Contents</a>
#             </p>
#             """

#             if added_anything:
#                 combined_html = "".join(html_parts) + footer_links
#                 tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
#                 self._html_to_pdf(combined_html, tmp.name)
#                 self.merger.append(tmp.name)
#                 temp_files.append(tmp.name)

#                 for pdf in pdf_files:
#                     self.merger.append(pdf)

#                 toc_rows.append(f"""
#                     <tr>
#                        <td>{i}</td>
#                        <td>{bank_code}</td>
#                        <td>{bank_type}</td>
#                        <td><a href='#{bank_anchor}'>{bank_name}</a></td>
#                     </tr>
#                 """)

#         # ===== Table of Contents (real table) =====
#         toc_html = f"""
#         {toc_anchor}
#         <h1 style='text-align:center;color:#1f4e79;'>Table of Contents</h1>
#         <p style='text-align:center;font-size:10pt;color:#666;'>
#            Generated on {datetime.now().strftime("%d %b %Y")}
#         </p>
#         <table border='1' cellspacing='0' cellpadding='4' 
#                style='width:90%;margin:auto;border-collapse:collapse;
#                       font-size:9pt;'>
#            <thead style='background:#f0f0f0;'>
#               <tr>
#                  <th>Sr No</th>
#                  <th>Code</th>
#                  <th>Bank Type</th>
#                  <th>Bank Name (Click to Open)</th>
#               </tr>
#            </thead>
#            <tbody>
#               {''.join(toc_rows)}
#            </tbody>
#         </table>
#         """
#         toc_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
#         self._html_to_pdf(toc_html, toc_pdf.name)
#         self.merger.merge(position=0, fileobj=toc_pdf.name)
#         toc_pdf.close()

#         # ===== finalize =====
#         self.merger.write(self.output_path)
#         self.merger.close()

#         reader = PdfReader(self.output_path)
#         writer = PdfWriter()
#         for page in reader.pages:
#             writer.add_page(page)

#         for i, title in enumerate(self.bank_titles, start=1):
#             writer.add_outline_item(title, i)
#         with open(self.output_path, "wb") as f:
#             writer.write(f)

#         print(f"✅ Final Linked PDF report generated at {self.output_path}")